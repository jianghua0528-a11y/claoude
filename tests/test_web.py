import json
from datetime import date

import pytest
from fastapi.testclient import TestClient

from cgroup.web.app import app
from cgroup.db.models import (ReviewItem, Order, Advance, MamaSettlement, Artist, Mama,
                              MamaAssistant, OperationLog)

AUTH = ("admin", "test-pw")


@pytest.fixture
def c():
    with TestClient(app) as client:
        client.auth = AUTH
        yield client


def _review(s, payload, group="报单群", raw="原文"):
    it = ReviewItem(source_group=group, raw_message=raw, status="待审",
                    parsed_json=json.dumps(payload, ensure_ascii=False) if payload else None)
    s.add(it); s.commit()
    return it.id


def test_auth_required(c):
    assert TestClient(app).get("/").status_code == 401
    assert TestClient(app).get("/", auth=("admin", "wrong")).status_code == 401
    assert TestClient(app).get("/", auth=("hacker", "test-pw")).status_code == 401
    assert c.get("/").status_code == 200


def test_every_page_renders(c, seeded):
    _review(seeded, {"艺人": "桃子", "妈咪": "小宝", "K": 1000, "日期": "2026-06-01"})
    c.post("/review/1/confirm")
    for url in ["/", "/?m=2026-06", "/?m=all", "/review", "/review?show=已确认", "/review/1", "/orders",
                "/orders?q=桃&st=全部&credit=due", "/orders/1/edit", "/upload", "/logs",
                "/finance", "/finance?tab=expense", "/finance?tab=settle", "/finance?tab=rate",
                "/finance/new?type=advance", "/dicts", "/dicts?kind=mama", "/dicts?kind=venue",
                "/dicts?kind=broker", "/dicts/artist/new", "/dicts/mama/1", "/health"]:
        r = c.get(url)
        assert r.status_code == 200, url


def test_confirm_is_idempotent(c, seeded):
    rid = _review(seeded, {"艺人": "peach", "妈咪": "阿美", "K": 3000, "日期": "2026-06-16"})
    c.post(f"/review/{rid}/confirm")
    c.post(f"/review/{rid}/confirm")                # 双击 / 重复提交
    seeded.expire_all()
    assert seeded.query(Order).count() == 1
    assert seeded.get(ReviewItem, rid).status == "已确认"


def test_confirm_bad_payload_shows_error_not_crash(c, seeded):
    rid = _review(seeded, {"艺人": "不存在的人", "K": 100})
    r = c.post(f"/review/{rid}/confirm")
    assert r.status_code == 200 and "不在字典" in r.text
    seeded.expire_all()
    assert seeded.query(Order).count() == 0 and seeded.get(ReviewItem, rid).status == "待审"


def test_edit_before_confirm(c, seeded):
    rid = _review(seeded, {"艺人": "不存在的人", "K": 100})
    form = {"edited": "1", "艺人": "桃子", "妈咪": "", "日期": "2026-06-20", "K": "1200", "M": "", "O": "0",
            "合作模式": "", "流向": ""}
    c.post(f"/review/{rid}/confirm", data=form)
    seeded.expire_all()
    o = seeded.query(Order).one()
    assert o.credit_k == 1200 and o.mode == "自单" and o.biz_date == date(2026, 6, 20)


def test_html_is_escaped(c, seeded):
    _review(seeded, None, group="财务群", raw="<script>alert(1)</script>")
    r = c.get("/review")
    assert "<script>alert(1)</script>" not in r.text and "&lt;script&gt;" in r.text
    assert "/finance/new?review=1" in r.text


def test_finance_review_to_advance_and_payslip_deduction(c, seeded):
    rid = _review(seeded, None, group="财务群", raw="桃子 预支 500")
    assert "预支/罚款" in c.get(f"/finance/new?review={rid}").text
    c.post("/finance/new", data={"type": "advance", "review": str(rid), "subject_type": "艺人",
                                 "subject_name": "peach", "kind": "预支", "amount": "500",
                                 "currency": "MYR", "related_month": "2026-06", "notes": ""})
    seeded.expire_all()
    a = seeded.query(Advance).one()
    assert a.subject_name == "桃子" and a.status == "待扣"
    assert seeded.get(ReviewItem, rid).status == "已确认"
    c.post("/finance/advance/settle_month?m=2026-06")
    seeded.expire_all()
    assert seeded.query(Advance).one().status == "已扣"


def test_finance_bad_input_rerenders(c, seeded):
    r = c.post("/finance/new", data={"type": "expense", "spend_date": "2026-06-01", "category": "伙食",
                                     "amount": "abc"})
    assert r.status_code == 200 and "格式不对" in r.text


def test_mama_settle_all(c, seeded):
    for k in (1000, 2000):
        _review(seeded, {"艺人": "桃子", "妈咪": "小宝", "K": k, "O": 150, "日期": "2026-06-01"})
    c.post("/review/1/confirm"); c.post("/review/2/confirm")
    mid = seeded.query(Mama).one().id
    c.post(f"/mama/{mid}/settle_all")
    seeded.expire_all()
    assert all(o.credit_status == "已结" and o.credit_ref for o in seeded.query(Order))
    st = seeded.query(MamaSettlement).one()
    assert st.amount == pytest.approx(3000 * 0.8 + 300)


def test_order_edit_void_restore(c, seeded):
    _review(seeded, {"艺人": "桃子", "妈咪": "小宝", "K": 1000, "日期": "2026-06-01"})
    c.post("/review/1/confirm")
    oid = seeded.query(Order).one().id
    c.post(f"/orders/{oid}/credit_toggle")
    r = c.post(f"/orders/{oid}/edit", data={"artist": "小白", "mama": "", "mode": "标准", "credit_k": "1"},
               follow_redirects=False)
    assert "err=" in r.headers["location"]            # 标准档没妈咪 → 拒
    c.post(f"/orders/{oid}/edit", data={"biz_date": "2026-06-02", "artist": "小白", "mama": "宝姐",
                                        "mode": "标准", "credit_k": "1500", "cash_m": "0", "ticket_o": "0"})
    seeded.expire_all()
    o = seeded.get(Order, oid)
    assert o.credit_k == 1500 and o.credit_status == "已结"     # 改金额保留已结
    c.post(f"/orders/{oid}/void")
    seeded.expire_all()
    assert seeded.get(Order, oid).status == "作废" and seeded.get(Order, oid).credit_status is None
    c.post(f"/orders/{oid}/restore")
    seeded.expire_all()
    assert seeded.get(Order, oid).credit_status == "已结"
    assert seeded.query(OperationLog).filter_by(operator="admin").count() >= 4


def test_export_xlsx(c, seeded):
    _review(seeded, {"艺人": "桃子", "K": 1000, "日期": "2026-06-01"})
    c.post("/review/1/confirm")
    r = c.get("/export?m=2026-06")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    from openpyxl import load_workbook
    import io
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 2 and ws.cell(row=2, column=3).value == "桃子"


def test_dicts_crud(c, seeded):
    c.post("/dicts/artist/new", data={"name": "新人", "aliases": "xr", "broker_id": "", "active": "1"})
    seeded.expire_all()
    assert seeded.query(Artist).filter_by(name="新人").one().aliases == "xr"
    r = c.post("/dicts/artist/new", data={"name": "桃子"})
    assert "已经有" in r.text
    mid = seeded.query(Mama).one().id
    c.post(f"/dicts/mama/{mid}", data={"name": "小宝", "assistants": "阿美，小红", "settlement_cycle": "月结"})
    seeded.expire_all()
    assert sorted(a.name for a in seeded.query(MamaAssistant)) == ["小红", "阿美"]
    c.post(f"/dicts/mama/{mid}", data={"name": "小宝", "assistants": "小红", "settlement_cycle": "月结"})
    seeded.expire_all()
    assert [a.name for a in seeded.query(MamaAssistant)] == ["小红"]
    c.post("/dicts/broker/new", data={"name": "老王", "pct": "10"})
    from cgroup.db.models import Broker
    seeded.expire_all()
    assert seeded.query(Broker).filter_by(name="老王").one().pct == pytest.approx(0.10)
